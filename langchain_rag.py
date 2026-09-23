from langchain_huggingface import HuggingFaceEmbeddings
from langchain_community.vectorstores import FAISS
from langchain_core.prompts import PromptTemplate
from langchain_deepseek import ChatDeepSeek
from langchain_core.runnables import (RunnableLambda,RunnablePassthrough)
from langchain_core.output_parsers import StrOutputParser
from pathlib import Path
from dotenv import load_dotenv
from docx import Document as DocxDocument
from langchain_core.documents import Document
import hashlib
import os

class LangChainRag:
    def __init__(self):
        base_path=Path(__file__).resolve().parent
        self.model_name="BAAI/bge-small-zh-v1.5"
        self.chat_model_name="deepseek-flash"

        self.file_path=(
            base_path
            /"rag_text"
            /"程序设计实训完整报告.docx"
        )

        self.vector_path=(
            base_path
            /"rag_text"
            /"langchain_faiss"
        )

        self.hash_path=(
            base_path
            /"rag_text"
            /"langchain_faiss.hash"
        )

        self.chunk_size=500
        self.chunk_strategy="question_title_v1"

        self.embeddings=None
        self.vector_store=None
        self.retriever=None
        self.chain=None

        self.question_titles = {
            "安全指数",
            "如此编码",
            "瑞瑞木板",
            "序列查询",
            "邻域均值",
            "相反数",
            "稀疏向量",
            "风险人群筛查",
            "学生排队",
            "消除类游戏",
            "两数之和",
            "有效的括号",
            "买卖股票的最佳时机",
            "爬楼梯",
            "最大子数组和"
        }

    def load(self):
        if self.vector_store is not None:
            return 
      
        current_hash=self.get_cache_hash()
        cache_valid=False

        if self.vector_path.exists() and self.hash_path.exists():
            old_hash=self.hash_path.read_text(
                encoding="utf-8"
            ).strip()
            if old_hash==current_hash:
                cache_valid=True
        if cache_valid:
            embeddings=self.get_embeddings()
            self.vector_store=FAISS.load_local(
                str(self.vector_path),
                embeddings,
                allow_dangerous_deserialization=True
            )
            return

        chunks=self.build_documents()
        embeddings=self.get_embeddings()
        self.vector_store=FAISS.from_documents(
            chunks,
            embeddings
        )

        self.vector_store.save_local(
            str(self.vector_path)
        )
        self.hash_path.write_text(
            current_hash,
            encoding="utf-8"
        )

    def build_chain(self):
        if self.retriever is None:
            self.build_retriever()

        prompt_template = PromptTemplate.from_template(
        """
        请只根据下面提供的资料回答问题。
        如果资料中没有足够信息，请明确说明。

        资料：
        {context}

        问题：
        {question}
        """
            )
        
        load_dotenv()
        model=ChatDeepSeek(
            model=self.chat_model_name,
            api_key=os.getenv("DEEPSEEK_API_KEY"),
            temperature=0
        )

        self.chain=(
            {
                "context":
                    self.retriever|RunnableLambda(self.format_documents),
                "question":RunnablePassthrough()
            }
            |prompt_template
            |model
            |StrOutputParser()
        )
    
    def build_retriever(self):
        if self.vector_store is None:
            self.load()
        self.retriever=self.vector_store.as_retriever(
            search_kwargs={
                "k":2
            }
        )


    def retrieve(self,question):#返回rag检索结果
        if self.retriever is None:
            self.build_retriever()
        return self.retriever.invoke(question)

    def ask(self,question):#检索加llm生成输出
        if self.chain is None:
            self.build_chain()
        return self.chain.invoke(question)
    
    def get_cache_hash(self):
        file_bytes=self.file_path.read_bytes()

        titles_text="\n".join(sorted(self.question_titles))
        config_text=(
            self.model_name
            +"\n"
            +str(self.chunk_size)
            +"\n"
            +str(titles_text)
            +"\n"
            +str(self.chunk_strategy)
        )
        hasher=hashlib.md5()
        hasher.update(file_bytes)#这个update函数不能接受字符串,接受的是字节数据
        hasher.update(config_text.encode("utf-8"))
        return hasher.hexdigest()
    
    def format_documents(self,documents):
        context_parts=[]
        for document in documents:
            context_parts.append(
                document.page_content
            )
        return "\n\n".join(context_parts)
    
    def get_embeddings(self):
        if self.embeddings is None:
            self.embeddings=HuggingFaceEmbeddings(
                model_name=self.model_name,
                model_kwargs={
                    "local_files_only":True
                }
            )
        return self.embeddings
    
    def build_documents(self):
        document=DocxDocument(self.file_path)
        texts=[]
        for paragraph in document.paragraphs:
            text=paragraph.text.strip()
            if text:
                texts.append(text)
            
        chunks=[]
        current_length=0
        current_parts=[]
        current_title=None

        for text in texts:
            if text in self.question_titles:
                if current_parts:
                    chunk="\n".join(current_parts)
                    chunks.append(
                        Document(
                            page_content=chunk,
                            metadata={
                                "source":self.file_path.name,
                                "section":current_title
                            }
                        )
                    )
                
                current_title=text 
                current_parts=[text]
                current_length=len(text)
                continue
            if current_parts and len(text)+current_length>self.chunk_size:
                chunk="\n".join(current_parts)
                chunks.append(
                        Document(
                            page_content=chunk,
                            metadata={
                                "source":self.file_path.name,
                                "section":current_title
                            }
                        )
                    )

                current_length=0
                current_parts=[]

                if current_title is not None:
                    current_parts.append(current_title)
                    current_length+=len(current_title)
            
            current_parts.append(text)
            current_length+=len(text)
        if current_parts:
            chunk="\n".join(current_parts)
            chunks.append(
                        Document(
                            page_content=chunk,
                            metadata={
                                "source":self.file_path.name,
                                "section":current_title
                            }
                        )
                    )
        return chunks
