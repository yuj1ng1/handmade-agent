from pathlib import Path

import numpy as np
from docx import Document
from sentence_transformers import SentenceTransformer


class KnowledgeBase:
    def __init__(self):
        self.model = None
        self.chunks = None
        self.embeddings = None

    def split_text(self, texts, chunk_size=500):
        chunks = []

        current_parts = []
        current_length = 0

        for text in texts:
            if (
                current_parts
                and current_length + len(text) > chunk_size
            ):
                chunk = "\n".join(current_parts)
                chunks.append(chunk)

                current_parts = []
                current_length = 0

            current_parts.append(text)
            current_length += len(text)

        if current_parts:
            chunk = "\n".join(current_parts)
            chunks.append(chunk)

        return chunks

    def load(self):
        if (
            self.model is not None
            and self.chunks is not None
            and self.embeddings is not None
        ):
            print("知识库已经加载，直接复用")
            return

        file_path = (
            Path(__file__).resolve().parent
            / "rag_text"
            / "程序设计实训完整报告.docx"
        )

        print("Word路径：", file_path)
        print("Word是否存在：", file_path.exists())

        if not file_path.exists():
            raise FileNotFoundError(
                f"找不到Word文件：{file_path}"
            )

        document = Document(file_path)

        print("普通段落数：", len(document.paragraphs))
        print("表格数：", len(document.tables))

        texts = []

        # 读取Word普通段落
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                texts.append(text)

        # 读取Word表格中的段落
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        text = paragraph.text.strip()

                        if text:
                            texts.append(text)

        print("读取到的文本段数：", len(texts))

        if texts:
            print("\n===== 前5个文本段 =====")

            for i, text in enumerate(texts[:5]):
                print(f"\n文本段 {i}")
                print(text)

        self.chunks = self.split_text(
            texts,
            chunk_size=500
        )

        print("\n生成的chunk数量：", len(self.chunks))

        if not self.chunks:
            raise ValueError(
                "没有生成任何chunk，说明Word内容没有正常读取"
            )

        print("\n===== 前3个chunk =====")

        for i, chunk in enumerate(self.chunks[:3]):
            print(f"\n----- chunk {i} -----")
            print(chunk)

        print("\n开始加载Embedding模型")

        self.model = SentenceTransformer(
            "BAAI/bge-small-zh-v1.5",
            local_files_only=True
        )

        embedding_path = (
            Path(__file__).resolve().parent
            / "rag_text"
            / "embedding.npy"
        )

        print("Embedding路径：", embedding_path)

        if embedding_path.exists():
            loaded_embeddings = np.load(
                embedding_path
            )

            print(
                "旧embedding shape：",
                loaded_embeddings.shape
            )

            # 检查缓存数量是否和当前chunk一致
            if (
                len(loaded_embeddings) == len(self.chunks)
                and len(loaded_embeddings) > 0
            ):
                self.embeddings = loaded_embeddings

                print("旧embedding有效，直接使用")

            else:
                print(
                    "旧embedding和当前chunk不匹配，重新生成"
                )

                self.embeddings = self.model.encode(
                    self.chunks
                )

                np.save(
                    embedding_path,
                    self.embeddings
                )

        else:
            print(
                "没有embedding.npy，首次生成"
            )

            self.embeddings = self.model.encode(
                self.chunks
            )

            np.save(
                embedding_path,
                self.embeddings
            )

        print(
            "最终embedding shape：",
            self.embeddings.shape
        )

    def search(self, question: str, top_k=3):
        self.load()

        question_embedding = self.model.encode(
            question
        )

        scores = []

        for i, embedding in enumerate(self.embeddings):
            score = (
                np.dot(
                    question_embedding,
                    embedding
                )
                /
                (
                    np.linalg.norm(question_embedding)
                    * np.linalg.norm(embedding)
                )
            )

            scores.append(
                (i, score)
            )

        scores.sort(
            key=lambda x: x[1],
            reverse=True
        )

        results = []

        for i, similarity in scores[:top_k]:
            results.append({
                "chunk_index": i,
                "similarity": float(similarity),
                "text": self.chunks[i]
            })

        return results


if __name__ == "__main__":
    knowledge_base = KnowledgeBase()

    print(
        "\n========== 测试知识库加载 ==========\n"
    )

    knowledge_base.load()

    print(
        "\n========== 测试知识库检索 ==========\n"
    )

    results = knowledge_base.search(
        "稀疏向量"
    )

    for result in results:
        print(
            "\n================================"
        )

        print(
            "chunk_index：",
            result["chunk_index"]
        )

        print(
            "similarity：",
            result["similarity"]
        )

        print("text：")
        print(result["text"])